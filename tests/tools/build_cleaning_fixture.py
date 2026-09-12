#!/usr/bin/env python3
"""Build the synthetic DOCX used by cleaning/preflight acceptance tests.

The output is intentionally dirty and must never be submitted. It contains a
comment, tracked replacement, draft watermark, internal ID, unresolved
placeholder, hidden text, and identifying core metadata.
"""

from __future__ import annotations

import argparse
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_east_asian_font(run, name="宋体"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def build_base(path: Path, cleanable: bool = False):
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run("TEST-ONLY 清洁流程测试文书")
    set_east_asian_font(run, "黑体")
    run.bold = True
    run.font.size = Pt(18)

    warning = document.add_paragraph()
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = warning.add_run("本文件完全虚构，仅用于离线测试，不得提交。")
    set_east_asian_font(run)
    run.bold = True
    run.font.color.rgb = RGBColor(0xC2, 0x41, 0x0C)

    paragraphs = [
        "内部编号：INT-TEST-CASE-001",
        "法院候选段落：原始争议金额为人民币10,000元。",
        "被告尚未付款。",
        "[内部备注：核对日期后再决定是否引用。]",
    ]
    if not cleanable:
        paragraphs.append("[[PENDING:确认送达地址]]")
    else:
        paragraphs.append("送达信息：TEST-ONLY 虚构地址（已确认）。")
    for text in paragraphs:
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
        paragraph.paragraph_format.line_spacing = 1.5

    if not cleanable:
        hidden = document.add_paragraph().add_run("INTERNAL-HIDDEN-TEXT")
        set_east_asian_font(hidden)
        hidden.font.hidden = True

    document.core_properties.title = "TEST-ONLY dirty legal document fixture"
    document.core_properties.author = "TEST-ONLY Internal Reviewer"
    document.core_properties.last_modified_by = "TEST-ONLY Drafting Agent"
    document.core_properties.keywords = "TEST-ONLY,INTERNAL,DRAFT"
    document.core_properties.comments = "Internal metadata must be removed from an authorized derived copy."
    document.save(path)


def run_helper(python: Path, helper: Path, *args: str):
    subprocess.run([str(python), str(helper), *args], check=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--python", type=Path, required=True)
    parser.add_argument("--document-skill-dir", type=Path, required=True)
    parser.add_argument("--cleanable", action="store_true", help="omit unresolved placeholder and hidden text")
    args = parser.parse_args()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    base = args.output.with_name("_base.docx")
    commented = args.output.with_name("_comments.docx")
    tracked = args.output.with_name("_tracked.docx")
    build_base(base, cleanable=args.cleanable)

    scripts = args.document_skill_dir / "scripts"
    run_helper(
        args.python,
        scripts / "comments_add.py",
        str(base),
        "--out",
        str(commented),
        "--author",
        "TEST-ONLY Reviewer",
        "--add",
        "法院候选段落=TEST-ONLY 内部批注：复核来源定位。",
        "--require_all",
    )
    run_helper(
        args.python,
        scripts / "add_tracked_replacements.py",
        str(commented),
        "--out",
        str(tracked),
        "--author",
        "TEST-ONLY Reviewer",
        "--replace",
        "被告尚未付款。=被告在约定期限内未付款。",
    )
    run_helper(
        args.python,
        scripts / "watermark_add.py",
        str(tracked),
        "--out",
        str(args.output),
        "--text",
        "DRAFT-TEST-ONLY",
    )

    for stage in (base, commented, tracked):
        stage.unlink(missing_ok=True)
    print(args.output)


if __name__ == "__main__":
    main()
