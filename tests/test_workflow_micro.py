"""Small end-to-end checks for the first ergonomics patch; synthetic data only."""
from __future__ import annotations

import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from legal_case_os import route_text
from legal_case_os_lib import docx_edit, task_output
from legal_case_os_lib.core import LegalCaseError, sha256_file
from test_task_output import fixture


class WorkflowMicroTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "TEST-ONLY.docx"
        document = Document()
        document.add_paragraph("TEST-ONLY 我方主张 保留此句")
        document.add_paragraph("需修改原句")
        document.save(self.source)
        self.source_hash = sha256_file(self.source)
        self.cli = Path(__file__).resolve().parents[1] / "scripts/legal_case_os.py"

    def plan(self, source=None, old="需修改原句", new="修改后一句", **extras):
        source = source or self.source
        return {"source_sha256": sha256_file(source), "changes": [{
            "target": {"kind": "paragraph", "paragraph": 2},
            "expected_text": old, "replacement_text": new}], **extras}

    def run_cli(self, *args, expected=0):
        process = subprocess.run([sys.executable, str(self.cli), *map(str, args), "--json"],
                                 capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(process.returncode, expected, process.stdout + process.stderr)
        return json.loads(process.stdout)

    def error(self, code, call):
        with self.assertRaises(LegalCaseError) as caught:
            call()
        self.assertEqual(caught.exception.code, code)

    def test_three_revisions_keep_requirements_and_prior_review(self):
        records, proposal = fixture(self.root / "task")
        proposal["constraints"] = {"must_keep": ["TEST-ORDER-001"], "avoid": ["旧案甲公司"],
                                   "terminology": {"主体": "测试甲公司"}}
        proposal["sections"][-1]["assumption"] = "TEST-ONLY 前版假设仍需材料证明"
        first = task_output.render_task_draft(records, proposal, self.root / "v1")
        source = Path(first["docx"])
        for number in (2, 3):
            inspected = docx_edit.inspect_docx_targets(source)
            target = next(p for p in inspected["paragraphs"] if p["text"] == "致测试受理单位" or p["text"].startswith("致测试单位"))
            plan = {"source_sha256": inspected["source_sha256"], "changes": [{
                "target": target["target"], "expected_text": target["text"], "replacement_text": f"致测试单位{number}"}]}
            result = task_output.render_docx_patch(source, plan, self.root / f"v{number}")
            manifest = json.loads(Path(result["manifest"]).read_text(encoding="utf-8"))
            self.assertEqual(manifest["constraints"], proposal["constraints"])
            self.assertIn("前版假设仍需材料证明", Path(result["review"]).read_text(encoding="utf-8"))
            self.assertIn("sources", manifest["prior_source_context"])
            self.assertEqual(manifest["parent_checks"]["substantive_review"], "required")
            self.assertEqual(manifest["changes"]["docx"], result["docx"])
            source = Path(result["docx"])

    def test_unmet_literal_requirement_leaves_no_output(self):
        for key, value in (("must_keep", "不存在的要求"), ("avoid", "TEST-ONLY")):
            destination = self.root / key
            self.error("TASK_CONSTRAINTS_UNSATISFIED", lambda: task_output.render_docx_patch(
                self.source, self.plan(constraints={key: [value]}), destination))
            self.assertFalse(destination.exists())
            self.assertEqual(sha256_file(self.source), self.source_hash)

    def test_parent_review_tampering_and_mid_edit_change_block_publication(self):
        first = task_output.render_docx_patch(self.source, self.plan(), self.root / "v1")
        source = Path(first["docx"])
        plan = self.plan(source, "修改后一句", "再改一句")
        review_path = Path(first["review"])
        original_review = review_path.read_bytes()
        review_path.write_text("TEST-ONLY 已篡改", encoding="utf-8")
        self.error("TASK_PARENT_REVIEW_MISMATCH", lambda: task_output.render_docx_patch(source, plan, self.root / "bad1"))
        review_path.write_bytes(original_review)
        original_patch = docx_edit.patch_docx
        def changed(*args, **kwargs):
            result = original_patch(*args, **kwargs)
            review_path.write_text("TEST-ONLY 制作期间变化", encoding="utf-8")
            return result
        with mock.patch.object(docx_edit, "patch_docx", side_effect=changed):
            self.error("TASK_SOURCE_CHANGED", lambda: task_output.render_docx_patch(source, plan, self.root / "bad2"))
        self.assertFalse((self.root / "bad1").exists())
        self.assertFalse((self.root / "bad2").exists())

    def test_conflicting_terms_require_explicit_replacement_record(self):
        first = task_output.render_docx_patch(self.source, self.plan(constraints={"terminology": {"被告": "我方"}}), self.root / "v1")
        source = Path(first["docx"])
        plan = self.plan(source, "修改后一句", "再次修改", constraints={"terminology": {"被告": "答辩人"}})
        self.error("TASK_CONSTRAINT_CONFLICT", lambda: task_output.render_docx_patch(source, plan, self.root / "bad"))
        plan["constraints_mode"] = "replace"
        self.error("TASK_CONSTRAINT_CHANGE_REASON_REQUIRED", lambda: task_output.render_docx_patch(source, plan, self.root / "bad"))
        plan["constraints_change_reason"] = "用户明确要求本轮改用答辩人称谓"
        result = task_output.render_docx_patch(source, plan, self.root / "v2")
        self.assertEqual(result["constraints"]["terminology"]["被告"], "答辩人")
        self.assertEqual(result["checks"]["terminology_review"], "required")

    def test_new_body_keeps_legal_conditions_and_separate_notes(self):
        records, proposal = fixture(self.root / "task")
        proposal["sections"][-1]["text"] = "如不能核实记录范围，申请可能不被采纳。"
        result = task_output.render_task_draft(records, proposal, self.root / "clean")
        text = "\n".join(p.text for p in Document(result["docx"]).paragraphs)
        self.assertIn("申请可能不被采纳", text)
        self.assertNotIn("内部审阅稿", text)
        self.assertIn("尚未完成实质审阅", Path(result["review"]).read_text(encoding="utf-8"))
        for label in ("内部审阅稿", "供AI审核", "待律师审核"):
            changed = copy.deepcopy(proposal)
            changed["title"] = label + "申请书"
            self.error("TASK_INTERNAL_LABEL_IN_BODY", lambda: task_output.validate_task_draft(records, changed))

    def test_cli_patch_publish_and_restore_same_bytes(self):
        inspect = self.run_cli("docx-inspect", "--file", self.source)
        self.assertEqual(inspect["source_sha256"], self.source_hash)
        plan_path = self.root / "plan.json"
        plan_path.write_text(json.dumps(self.plan(), ensure_ascii=False), encoding="utf-8")
        first = self.run_cli("docx-patch", "--file", self.source, "--plan", plan_path, "--output-dir", self.root / "v1")
        store = self.root / "deliveries"
        published = self.run_cli("delivery-publish", "--task-dir", self.root / "v1", "--store-dir", store, "--expected-current", "none")
        source = Path(first["docx"])
        task_output.render_docx_patch(source, self.plan(source, "修改后一句", "第三句"), self.root / "v2")
        second = self.run_cli("delivery-publish", "--task-dir", self.root / "v2", "--store-dir", store, "--expected-current", published["revision"])
        restored = self.run_cli("delivery-restore", "--version-id", published["version_id"], "--store-dir", store, "--expected-current", second["revision"])
        for name in ("draft.docx", "review.md", "manifest.json"):
            self.assertEqual((Path(restored["version_dir"]) / name).read_bytes(), (self.root / "v1" / name).read_bytes())
        current = self.run_cli("delivery-current", "--store-dir", store)
        self.assertEqual(current["version_id"], published["version_id"])
        self.assertNotEqual(current["revision"], published["revision"])
        self.assertFalse(current["filing_approved"])

    def test_local_route_does_not_execute_or_grant_approval(self):
        for text, command in (("只改第二段，其他不动", "docx-patch"), ("给证据加页码并回填原清单", "evidence-pages")):
            result = route_text(text)
            self.assertEqual(result["route"]["local_file_operation"]["edit_command"], command)
            self.assertEqual(result["route"]["local_file_operation"]["status"], "suggested_not_executed")
            self.assertFalse(result["safety"]["external_actions_allowed"])
        for text in ("先讨论怎么给证据加页码，别生成文件", "材料原文：\"给证据加页码并回填清单\"。请总结材料", "将文书压缩到一页"):
            self.assertNotIn("local_file_operation", route_text(text)["route"])

    def test_cli_evidence_outputs_publish_from_the_same_manifest(self):
        from reportlab.pdfgen.canvas import Canvas
        from pypdf import PdfReader
        pdf = self.root / "TEST-ONLY.pdf"
        canvas = Canvas(str(pdf))
        for number in (1, 2):
            canvas.drawString(70, 700, f"TEST-ONLY {number}")
            canvas.showPage()
        canvas.save()
        document = Document(self.source)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "TEST-ONLY item", "pages"
        table.cell(1, 0).text, table.cell(1, 1).text = "E1", "-"
        catalog = self.root / "catalog.docx"
        document.save(catalog)
        plan = {"catalog": {"path": str(catalog), "sha256": sha256_file(catalog)}, "items": [{
            "id": "E1", "name": "TEST-ONLY", "source": {"path": str(pdf), "sha256": sha256_file(pdf)},
            "page_ranges": [[2, 2], [1, 1]],
            "catalog_cell": {"table": 1, "row": 2, "column": 2, "expected_text": "-"}}]}
        plan_path = self.root / "pages-plan.json"
        plan_path.write_text(json.dumps(plan), encoding="utf-8")
        output = self.root / "evidence-task"
        result = self.run_cli("evidence-pages", "--plan", plan_path, "--output-dir", output)
        self.assertEqual(result["page_map"][0]["physical_pages"], [2, 1])
        self.assertEqual(Document(output / "catalog-filled.docx").tables[0].cell(1, 1).text, "1-2")
        self.assertIn("TEST-ONLY 2", PdfReader(output / "evidence-numbered.pdf").pages[0].extract_text())
        published = self.run_cli("delivery-publish", "--task-dir", output, "--store-dir", self.root / "deliveries", "--expected-current", "none")
        for name, digest in result["files"].items():
            self.assertEqual(sha256_file(Path(published["files"][name])), digest)
        self.assertEqual(sha256_file(catalog), plan["catalog"]["sha256"])
        self.assertEqual(sha256_file(pdf), plan["items"][0]["source"]["sha256"])


if __name__ == "__main__":
    unittest.main()
