"""TEST-ONLY integration: usable tables and source-bound learning in real drafts."""
from __future__ import annotations

import copy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from legal_case_os_lib import task_output
from legal_case_os_lib.core import LegalCaseError, sha256_file
from legal_case_os_lib.docx_edit import inspect_docx_targets
from legal_case_os_lib.learning import build_learning_candidate
from legal_case_os_lib.materials import ingest_material, read_material
from test_task_output import fixture


class FinalPolishTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="TEST-ONLY-polish-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.records, self.proposal = fixture(self.root / "inputs")
        source = self.root / "TEST-ONLY-method.txt"
        source.write_text("TEST-ONLY 先列明确请求，再说明根据，并写明条件。", encoding="utf-8")
        record = ingest_material(source, self.root / "learning-materials", "exemplar")
        text = read_material(record)["text"]
        quote = "先列明确请求，再说明根据，并写明条件。"
        start = text.index(quote)
        self.learning_proposal = {"id": "L-TEST-ONLY", "title": "TEST-ONLY 方法", "reasoning_patterns": [{
            "id": "R1", "text": "先界定请求再展开理由。", "applies_when": ["资料足够确定请求范围"],
            "not_applicable_when": ["仅填写固定身份栏"],
            "source_bindings": [{"id": record["id"], "start": start, "end": start + len(quote), "quote": quote}],
            "application_guide": {"steps": ["先定范围", "核对材料", "说明限制"],
                "good_example": "在能定位的期间内申请调取记录。", "bad_example": "无材料也断言全部成立。",
                "difference": "保留条件，不把旧案结果移入本案。"}}]}
        learning = build_learning_candidate([record], self.learning_proposal, self.root / "candidate.json")
        self.learning_source = source
        self.learning = Path(learning["candidate_path"])
        baseline = task_output.render_task_draft(self.records, self.proposal, self.root / "unbound")
        self.draft = Path(baseline["output_dir"]) / "draft.md"
        self.text = self.draft.read_text(encoding="utf-8")
        actual_quote = "申请范围以明确的记录标识和期间为限，便于核对材料。"
        offset = self.text.index(actual_quote)
        self.application = {"entries": [{"entry_id": "R1", "category": "reasoning_patterns", "source_binding_index": 1,
            "output_binding": {"start": offset, "end": offset + len(actual_quote), "quote": actual_quote},
            "adaptation": "只采用请求限定与理由对应的步骤，仍保留本任务资料范围。",
            "condition_checks": [
                {"field": "applies_when", "index": 1, "status": "unknown", "basis": "TEST-ONLY 尚需确认记录完整。"},
                {"field": "not_applicable_when", "index": 1, "status": "absent", "basis": "本稿有请求及理由。"}]}]}

    def with_application(self):
        proposal = copy.deepcopy(self.proposal)
        proposal["learning_applications"] = [{"learning_path": str(self.learning), "application": self.application}]
        return proposal

    def test_unknown_conditions_stay_in_review_and_bind_actual_markdown(self):
        result = task_output.render_task_draft(self.records, self.with_application(), self.root / "bound")
        manifest = json.loads(Path(result["manifest"]).read_text(encoding="utf-8"))
        report = manifest["learning_application_reports"][0]
        actual = (self.root / "bound/draft.md").read_text(encoding="utf-8")
        self.assertEqual(report["output_text_sha256"], hashlib.sha256(actual.encode("utf-8")).hexdigest())
        self.assertEqual(report["entries"][0]["application_status"], "needs_review")
        self.assertIn("尚需确认记录完整", Path(result["review"]).read_text(encoding="utf-8"))
        self.assertNotIn("尚需确认记录完整", "\n".join(p.text for p in Document(result["docx"]).paragraphs))
        self.assertFalse(report["checks"]["semantic_truth_verified_by_script"])
        self.assertFalse(report["persisted_to_library"])

    def test_wrong_actual_output_binding_never_publishes(self):
        self.application["entries"][0]["output_binding"]["quote"] = "TEST-ONLY 假引文"
        with self.assertRaises(LegalCaseError) as caught:
            task_output.render_task_draft(self.records, self.with_application(), self.root / "bad")
        self.assertEqual(caught.exception.code, "LEARNING_APPLICATION_QUOTE")
        self.assertFalse((self.root / "bad").exists())

    def test_learning_changes_during_render_stop_publication(self):
        original = task_output.create_docx_from_markdown
        def changed(*args, **kwargs):
            output = original(*args, **kwargs)
            self.learning.write_bytes(self.learning.read_bytes() + b"\n")
            return output
        with mock.patch.object(task_output, "create_docx_from_markdown", side_effect=changed):
            with self.assertRaises(LegalCaseError) as caught:
                task_output.render_task_draft(self.records, self.with_application(), self.root / "changed")
        self.assertEqual(caught.exception.code, "TASK_LEARNING_SOURCE_CHANGED")
        self.assertFalse((self.root / "changed").exists())

    def test_learning_source_changes_without_json_change_are_rejected(self):
        learning_hash = sha256_file(self.learning)
        original = task_output.create_docx_from_markdown
        def changed(*args, **kwargs):
            output = original(*args, **kwargs)
            self.learning_source.write_text("TEST-ONLY 来源已变", encoding="utf-8")
            return output
        with mock.patch.object(task_output, "create_docx_from_markdown", side_effect=changed):
            with self.assertRaises(LegalCaseError) as caught:
                task_output.render_task_draft(self.records, self.with_application(), self.root / "changed-source")
        self.assertIn(caught.exception.code, {"LEARNING_SOURCE_STALE", "LEARNING_SOURCE_INVALID"})
        self.assertEqual(sha256_file(self.learning), learning_hash)
        self.assertFalse((self.root / "changed-source").exists())

    def test_current_material_changes_during_render_are_rejected(self):
        original = task_output.create_docx_from_markdown
        def changed(*args, **kwargs):
            output = original(*args, **kwargs)
            Path(self.records[1]["text_path"]).write_text("TEST-ONLY 本次材料已变", encoding="utf-8")
            return output
        with mock.patch.object(task_output, "create_docx_from_markdown", side_effect=changed):
            with self.assertRaises(LegalCaseError) as caught:
                task_output.render_task_draft(self.records, self.with_application(), self.root / "changed-current")
        self.assertEqual(caught.exception.code, "TASK_SOURCE_CHANGED")
        self.assertFalse((self.root / "changed-current").exists())

    def test_local_revision_preserves_prior_application_but_requires_recheck(self):
        first = task_output.render_task_draft(self.records, self.with_application(), self.root / "v1")
        inspected = inspect_docx_targets(first["docx"])
        item = next(p for p in inspected["paragraphs"] if p["text"] == "致测试受理单位")
        plan = {"source_sha256": inspected["source_sha256"], "changes": [{
            "target": item["target"], "expected_text": item["text"], "replacement_text": "致另一测试单位"}]}
        second = task_output.render_docx_patch(first["docx"], plan, self.root / "v2")
        manifest = json.loads(Path(second["manifest"]).read_text(encoding="utf-8"))
        self.assertEqual(manifest["checks"]["learning_application_review"], "required")
        self.assertIn("learning_application_reports", manifest["prior_source_context"])
        self.assertIn("尚需确认记录完整", Path(second["review"]).read_text(encoding="utf-8"))

    def test_rebuilt_template_cannot_claim_original_layout_preservation(self):
        self.proposal["template_use"]["mode"] = "preserve"
        with self.assertRaises(LegalCaseError) as caught:
            task_output.validate_task_draft(self.records, self.proposal)
        self.assertEqual(caught.exception.code, "TASK_TEMPLATE_MODE_UNSUPPORTED")

    def test_table_in_task_output_is_real_word_table(self):
        self.proposal["sections"].append({"kind": "administrative", "text": "| TEST-ONLY 项目 | 页码 |\n| --- | ---: |\n| A | 1-2 |\n| B | 3 |"})
        result = task_output.render_task_draft(self.records, self.proposal, self.root / "table")
        document = Document(result["docx"])
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].cell(1, 1).text, "1-2")
        self.assertEqual(result["checks"]["template_layout"], "rebuilt_not_preserved")

    def test_cli_learning_check_without_current_fact_records(self):
        application_path = self.root / "application.json"
        application_path.write_text(json.dumps(self.application, ensure_ascii=False), encoding="utf-8")
        output = self.root / "application-report.json"
        cli = Path(__file__).resolve().parents[1] / "scripts/legal_case_os.py"
        process = subprocess.run([sys.executable, str(cli), "learning-check", "--learning", str(self.learning),
            "--application", str(application_path), "--draft-text", str(self.draft), "--output", str(output), "--json"],
            capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(process.returncode, 0, process.stdout + process.stderr)
        report = json.loads(output.read_text(encoding="utf-8"))
        self.assertEqual(report["draft_text_sha256"], sha256_file(self.draft))
        self.assertEqual(report["entries"][0]["application_status"], "needs_review")
        self.assertEqual(json.loads(process.stdout)["report_path"], str(output.resolve()))


if __name__ == "__main__":
    unittest.main()
