#!/usr/bin/env python3
"""Build and verify the TEST-ONLY v1.2 template boundary render suite.

The suite is deliberately self-contained: it creates one fictional DOCX form
fixture, treats that fixture as immutable after its first build, fills five
boundary scenarios by changing text nodes only, renders each derived DOCX to
PDF/PNG through the Windows short-path renderer, and records structural and
raster/PDF heuristics.  An optional ``visual-review.json`` may record an AI
page review, but it can never set lawyer approval or activate a template.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import stat
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import pdfplumber
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ROOT = PROJECT_ROOT / "tests" / "fixtures" / "template-boundary-v12"
RENDERER_CANDIDATES = (
    Path.home() / ".codex" / "skills" / "win-docx-render" / "scripts" / "render_docx_win.py",
    Path.home() / ".agents" / "skills" / "win-docx-render" / "scripts" / "render_docx_win.py",
)
SUITE_ID = "TEST-ONLY-template-boundary-v12"
FIXED_ZIP_TIME = (2026, 8, 25, 0, 0, 0)
PAGE_RE = re.compile(r"page_p(\d+)\.png$", re.IGNORECASE)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}

PLACEHOLDERS = {
    "party_name": "[[PARTY_NAME]]",
    "cause_of_action": "[[CAUSE_OF_ACTION]]",
    "court_name": "[[COURT_NAME]]",
    "amount": "[[AMOUNT]]",
    "identifier": "[[IDENTIFIER]]",
    "address": "[[ADDRESS]]",
    "optional_note": "[[OPTIONAL_NOTE]]",
}

BASE_VALUES = {
    "party_name": "泉州测试星河网络科技有限公司",
    "cause_of_action": "TEST-ONLY虚构买卖合同纠纷",
    "court_name": "福建省泉州市丰泽区人民法院",
    "amount": "人民币123,456.78元",
    "identifier": "91350500TESTONLY001",
    "address": "福建省泉州市丰泽区测试大道88号",
    "optional_note": "仅用于TEST-ONLY离线模板测试",
}

MAX_PARTY_NAME = (
    "泉州市测试星河超长主体名称技术研发供应链管理知识产权服务电子商务有限责任公司"
    "厦门测试分公司福州测试办事处（TEST-ONLY虚构边界主体）"
)
MAX_COURT_NAME = (
    "中华人民共和国福建省测试市测试区人民法院知识产权与涉外商事案件综合审判庭"
    "（TEST-ONLY虚构最长法院名称）"
)
LARGE_AMOUNT = "人民币9,999,999,999,999,999,999.99元"

SCENARIOS: tuple[dict[str, Any], ...] = (
    {
        "scenario_id": "normal-values",
        "label": "正常值",
        "values": dict(BASE_VALUES),
        "explicit_blank_fields": [],
    },
    {
        "scenario_id": "explicit-empty",
        "label": "明确空值",
        "values": {**BASE_VALUES, "optional_note": ""},
        "explicit_blank_fields": ["optional_note"],
    },
    {
        "scenario_id": "max-party-name",
        "label": "配置上限主体名称",
        "values": {**BASE_VALUES, "party_name": MAX_PARTY_NAME},
        "explicit_blank_fields": [],
    },
    {
        "scenario_id": "max-court-name",
        "label": "配置上限法院名称",
        "values": {**BASE_VALUES, "court_name": MAX_COURT_NAME},
        "explicit_blank_fields": [],
    },
    {
        "scenario_id": "large-amount",
        "label": "大额金额",
        "values": {**BASE_VALUES, "amount": LARGE_AMOUNT},
        "explicit_blank_fields": [],
    },
)


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_bytes_if_changed(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.is_file() and sha256_file(path) == sha256_bytes(data):
        return
    with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as handle:
        handle.write(data)
        temp_path = Path(handle.name)
    os.replace(temp_path, path)


def write_json_if_changed(path: Path, payload: Any) -> None:
    data = (json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")
    write_bytes_if_changed(path, data)


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def relpath(path: Path, root: Path) -> str:
    return os.path.relpath(path.resolve(), root.resolve()).replace("\\", "/")


def set_cell_margins(cell: Any, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table: Any, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def apply_font(run: Any, name: str, size: float, *, bold: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def build_source_docx(path: Path) -> None:
    document = Document()
    props = document.core_properties
    props.title = "TEST-ONLY template boundary fixture"
    props.subject = "Offline legal template rendering test"
    props.author = "legal-case-os"
    props.last_modified_by = "legal-case-os"
    props.created = datetime(2026, 8, 25, 0, 0, 0)
    props.modified = datetime(2026, 8, 25, 0, 0, 0)

    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    apply_font(title.add_run("TEST-ONLY 民事案件委托信息表"), "Microsoft YaHei", 16, bold=True, color="2E74B5")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(10)
    apply_font(subtitle.add_run("纯虚构模板边界渲染固定件｜不得用于真实案件"), "Microsoft YaHei", 10, color="555555")

    labels = (
        ("委托人名称", PLACEHOLDERS["party_name"]),
        ("案由", PLACEHOLDERS["cause_of_action"]),
        ("受理法院", PLACEHOLDERS["court_name"]),
        ("争议标的金额", PLACEHOLDERS["amount"]),
        ("主体识别码", PLACEHOLDERS["identifier"]),
        ("联系地址", PLACEHOLDERS["address"]),
        ("可选备注", PLACEHOLDERS["optional_note"]),
    )
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "字段"
    header[1].text = "拟填内容"
    for cell in header:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F4F7")
        cell._tc.get_or_add_tcPr().append(shading)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            apply_font(run, "Microsoft YaHei", 10.5, bold=True)

    for label, marker in labels:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = marker
        for index, cell in enumerate(cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            for run in paragraph.runs:
                apply_font(run, "Microsoft YaHei", 10.5, bold=index == 0)

    set_table_geometry(table, [2300, 7060])

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    apply_font(
        note.add_run("说明：空白字段仅在填写计划明确标注为空时保留；本固定件不包含律师授权或收费决定。"),
        "Microsoft YaHei",
        9,
        color="555555",
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    apply_font(footer.add_run("TEST-ONLY｜离线渲染测试｜禁止提交"), "Microsoft YaHei", 8, color="777777")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
        temp_docx = Path(handle.name)
    try:
        document.save(temp_docx)
        write_bytes_if_changed(path, canonical_docx_bytes(temp_docx))
    finally:
        temp_docx.unlink(missing_ok=True)


def canonical_docx_bytes(path: Path, overrides: dict[str, bytes] | None = None) -> bytes:
    overrides = overrides or {}
    with zipfile.ZipFile(path, "r") as source:
        members = {name: source.read(name) for name in source.namelist()}
    members.update(overrides)
    with tempfile.SpooledTemporaryFile() as output:
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
            for name in sorted(members):
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = 0o600 << 16
                archive.writestr(info, members[name])
        output.seek(0)
        return output.read()


def ensure_source(root: Path) -> tuple[Path, dict[str, Any]]:
    source_dir = root / "source"
    source_dir.mkdir(parents=True, exist_ok=True)
    source_path = source_dir / "TEST-ONLY-form-template.docx"
    baseline_path = source_dir / "source-baseline.json"
    if not source_path.exists():
        build_source_docx(source_path)
        try:
            source_path.chmod(stat.S_IREAD | stat.S_IRGRP | stat.S_IROTH)
        except OSError:
            pass
    source_hash = sha256_file(source_path)
    if baseline_path.exists():
        baseline = load_json(baseline_path)
        if baseline.get("sha256") != source_hash:
            raise RuntimeError("immutable TEST-ONLY source hash changed; refusing to rebuild outputs")
    else:
        baseline = {
            "test_only": True,
            "path": relpath(source_path, root),
            "sha256": source_hash,
            "created_by": "tools/build_template_boundary_suite.py",
            "immutable_after_first_build": True,
            "lawyer_approved_template": False,
        }
        write_json_if_changed(baseline_path, baseline)
    return source_path, baseline


def fill_docx_bytes(source: Path, values: dict[str, str]) -> bytes:
    with zipfile.ZipFile(source, "r") as archive:
        document_xml = archive.read("word/document.xml")
    for key, marker in PLACEHOLDERS.items():
        encoded_marker = marker.encode("utf-8")
        if document_xml.count(encoded_marker) != 1:
            raise RuntimeError(f"placeholder occurrence must be exactly one: {marker}")
        replacement = escape(values[key]).encode("utf-8")
        document_xml = document_xml.replace(encoded_marker, replacement)
    return canonical_docx_bytes(source, {"word/document.xml": document_xml})


def table_geometry_signature(path: Path) -> list[dict[str, Any]]:
    with zipfile.ZipFile(path, "r") as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    signatures: list[dict[str, Any]] = []
    for table in root.findall(".//w:tbl", NS):
        tbl_pr = table.find("w:tblPr", NS)
        tbl_grid = table.find("w:tblGrid", NS)

        def attrs(element: ET.Element | None) -> dict[str, str] | None:
            if element is None:
                return None
            return {key.split("}")[-1]: value for key, value in sorted(element.attrib.items())}

        rows = table.findall("w:tr", NS)
        signatures.append(
            {
                "table_width": attrs(tbl_pr.find("w:tblW", NS) if tbl_pr is not None else None),
                "table_indent": attrs(tbl_pr.find("w:tblInd", NS) if tbl_pr is not None else None),
                "layout": attrs(tbl_pr.find("w:tblLayout", NS) if tbl_pr is not None else None),
                "grid": [attrs(item) for item in (tbl_grid.findall("w:gridCol", NS) if tbl_grid is not None else [])],
                "row_count": len(rows),
                "rows": [
                    [
                        {
                            "width": attrs(cell.find("w:tcPr/w:tcW", NS)),
                            "grid_span": attrs(cell.find("w:tcPr/w:gridSpan", NS)),
                            "v_merge": attrs(cell.find("w:tcPr/w:vMerge", NS)),
                        }
                        for cell in row.findall("w:tc", NS)
                    ]
                    for row in rows
                ],
            }
        )
    return signatures


def non_document_parts_equal(source: Path, derived: Path) -> tuple[bool, list[str]]:
    with zipfile.ZipFile(source, "r") as left, zipfile.ZipFile(derived, "r") as right:
        left_names = set(left.namelist())
        right_names = set(right.namelist())
        changed = sorted(left_names ^ right_names)
        for name in sorted(left_names & right_names):
            if name == "word/document.xml":
                continue
            if left.read(name) != right.read(name):
                changed.append(name)
    return not changed, sorted(set(changed))


def document_text(path: Path) -> str:
    with zipfile.ZipFile(path, "r") as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    return "".join(element.text or "" for element in root.findall(".//w:t", NS))


def field_values(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path, "r") as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    output: dict[str, str] = {}
    for row in root.findall(".//w:tbl/w:tr", NS)[1:]:
        cells = row.findall("w:tc", NS)
        if len(cells) != 2:
            continue
        label = "".join(item.text or "" for item in cells[0].findall(".//w:t", NS))
        value = "".join(item.text or "" for item in cells[1].findall(".//w:t", NS))
        output[label] = value
    return output


def find_renderer(explicit: str | None) -> Path:
    candidates = ([Path(explicit)] if explicit else []) + list(RENDERER_CANDIDATES)
    for path in candidates:
        if path.is_file():
            return path.resolve()
    raise FileNotFoundError("win-docx-render helper was not found")


def page_files(render_dir: Path) -> list[Path]:
    numbered: list[tuple[int, Path]] = []
    for path in render_dir.glob("page_p*.png"):
        match = PAGE_RE.search(path.name)
        if match:
            numbered.append((int(match.group(1)), path))
    numbered.sort(key=lambda pair: pair[0])
    if numbered and [number for number, _ in numbered] != list(range(1, len(numbered) + 1)):
        raise RuntimeError(f"non-contiguous page set: {render_dir}")
    return [path for _, path in numbered]


def render_is_reusable(previous: dict[str, Any] | None, docx_hash: str, root: Path) -> bool:
    if not previous or previous.get("output_docx", {}).get("sha256") != docx_hash:
        return False
    render = previous.get("render", {})
    pdf = root / render.get("pdf", {}).get("path", "")
    if not pdf.is_file() or sha256_file(pdf) != render.get("pdf", {}).get("sha256"):
        return False
    pages = render.get("pages", [])
    if not pages:
        return False
    return all(
        (root / page["path"]).is_file() and sha256_file(root / page["path"]) == page["sha256"]
        for page in pages
    )


def render_docx(docx_path: Path, render_dir: Path, renderer: Path, *, skip_render: bool) -> tuple[Path, list[Path], str]:
    pdf_path = render_dir / "page.pdf"
    pages = page_files(render_dir)
    if skip_render:
        if not pdf_path.is_file() or not pages:
            raise RuntimeError(f"--skip-render requested but complete render is missing: {docx_path}")
        return pdf_path, pages, "reused_existing_render"

    render_dir.mkdir(parents=True, exist_ok=True)
    for stale in [pdf_path, *pages]:
        stale.unlink(missing_ok=True)
    command = [
        sys.executable,
        str(renderer),
        str(docx_path),
        "--outdir",
        str(render_dir),
        "--prefix",
        "page",
        "--dpi",
        "144",
    ]
    completed = subprocess.run(command, text=True, encoding="utf-8", errors="replace", capture_output=True, check=False)
    if completed.returncode:
        raise RuntimeError(f"render failed for {docx_path}:\n{completed.stdout}\n{completed.stderr}")
    pages = page_files(render_dir)
    if not pdf_path.is_file() or not pages:
        raise RuntimeError(f"renderer produced an incomplete PDF/PNG set: {docx_path}")
    return pdf_path, pages, "rendered_now"


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", value)


def detect_text_overlaps(chars: list[dict[str, Any]]) -> list[dict[str, Any]]:
    collisions: list[dict[str, Any]] = []
    visible = [item for item in chars if str(item.get("text", "")).strip()]
    for index, left in enumerate(visible):
        lx0, lx1 = float(left["x0"]), float(left["x1"])
        ltop, lbottom = float(left["top"]), float(left["bottom"])
        larea = max(0.01, (lx1 - lx0) * (lbottom - ltop))
        for right in visible[index + 1 :]:
            rx0, rx1 = float(right["x0"]), float(right["x1"])
            rtop, rbottom = float(right["top"]), float(right["bottom"])
            if rx0 > lx1 + 1 or rtop > lbottom + 1 or ltop > rbottom + 1:
                continue
            width = max(0.0, min(lx1, rx1) - max(lx0, rx0))
            height = max(0.0, min(lbottom, rbottom) - max(ltop, rtop))
            rarea = max(0.01, (rx1 - rx0) * (rbottom - rtop))
            ratio = (width * height) / min(larea, rarea)
            if ratio >= 0.65:
                collisions.append(
                    {
                        "left": str(left.get("text", "")),
                        "right": str(right.get("text", "")),
                        "overlap_ratio": round(ratio, 4),
                    }
                )
                if len(collisions) >= 20:
                    return collisions
    return collisions


def inspect_render(pdf_path: Path, pages: list[Path], values: dict[str, str]) -> dict[str, Any]:
    extracted_parts: list[str] = []
    content_stream_parts: list[str] = []
    collisions: list[dict[str, Any]] = []
    out_of_bounds_chars: list[dict[str, Any]] = []
    page_dimensions: list[dict[str, float]] = []
    with pdfplumber.open(pdf_path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            extracted_parts.append(page.extract_text() or "")
            # pdfplumber's layout extraction may interleave the left-column
            # label between two wrapped lines in the right cell.  The PDF
            # content-stream character order preserves the injected run and is
            # therefore the deterministic value-integrity signal here.
            content_stream_parts.append("".join(str(item.get("text", "")) for item in page.chars))
            page_dimensions.append({"width_pt": round(page.width, 3), "height_pt": round(page.height, 3)})
            for item in page.chars:
                if (
                    float(item["x0"]) < -0.25
                    or float(item["x1"]) > page.width + 0.25
                    or float(item["top"]) < -0.25
                    or float(item["bottom"]) > page.height + 0.25
                ):
                    out_of_bounds_chars.append(
                        {"page": page_number, "text": str(item.get("text", "")), "bbox": [item["x0"], item["top"], item["x1"], item["bottom"]]}
                    )
            collisions.extend({"page": page_number, **item} for item in detect_text_overlaps(page.chars))
        pdf_page_count = len(pdf.pages)

    extracted = "\n".join(extracted_parts)
    content_stream_text = "".join(content_stream_parts)
    normalized = normalize_text(content_stream_text)
    missing_values = [key for key, value in values.items() if value and normalize_text(value) not in normalized]
    replacement_glyphs = sorted({char for char in ("\ufffd", "�", "□") if char in content_stream_text})

    edge_results = []
    blank_pages = []
    for page_number, page_path in enumerate(pages, start=1):
        with Image.open(page_path) as image:
            gray = image.convert("L")
            width, height = gray.size
            pixels = gray.load()
            edge_count = 0
            nonwhite = 0
            total = width * height
            edge = 4
            for y in range(height):
                for x in range(width):
                    if pixels[x, y] < 245:
                        nonwhite += 1
                        if x < edge or x >= width - edge or y < edge or y >= height - edge:
                            edge_count += 1
            ratio = nonwhite / max(1, total)
            if ratio < 0.001:
                blank_pages.append(page_number)
            edge_results.append(
                {
                    "page": page_number,
                    "width_px": width,
                    "height_px": height,
                    "edge_nonwhite_pixels": edge_count,
                    "nonwhite_ratio": round(ratio, 8),
                }
            )

    checks = {
        "page_set_complete": {
            "status": "pass" if pdf_page_count == len(pages) and len(pages) > 0 else "fail",
            "method": "PDF page count equals contiguous PNG page count",
            "pdf_page_count": pdf_page_count,
            "png_page_count": len(pages),
        },
        "no_clipping": {
            "status": "pass" if not out_of_bounds_chars and all(item["edge_nonwhite_pixels"] == 0 for item in edge_results) else "fail",
            "method": "PDF character boxes stay within page bounds and 4-pixel raster edge bands contain no ink",
            "out_of_bounds_chars": out_of_bounds_chars[:20],
            "page_edge_metrics": edge_results,
        },
        "no_text_overlap": {
            "status": "pass" if not collisions else "fail",
            "method": "No pair of non-space PDF character boxes overlaps by 65% or more of the smaller glyph box",
            "collisions": collisions[:20],
            "heuristic_limit": "Does not prove absence of every graphical overlap; AI and lawyer visual review remain separate gates.",
        },
        "no_missing_glyphs": {
            "status": "pass" if not missing_values and not replacement_glyphs else "fail",
            "method": "All non-empty injected values are recoverable from PDF text and no replacement/tofu glyph is extracted",
            "missing_value_fields": missing_values,
            "replacement_glyphs": replacement_glyphs,
        },
        "no_blank_pages": {
            "status": "pass" if not blank_pages else "fail",
            "method": "Each raster page has at least 0.1% non-white pixels",
            "blank_pages": blank_pages,
        },
    }
    return {
        "page_dimensions": page_dimensions,
        "checks": checks,
        "all_render_checks_passed": all(item["status"] == "pass" for item in checks.values()),
    }


def load_visual_review(root: Path, scenario_records: list[dict[str, Any]]) -> dict[str, Any]:
    review_path = root / "visual-review.json"
    pending = {
        "status": "pending_ai_and_lawyer_visual_review",
        "review_path": relpath(review_path, root),
        "lawyer_approval": False,
        "lawyer_approval_substituted": False,
        "filing_use_allowed": False,
    }
    if not review_path.is_file():
        return pending
    review = load_json(review_path)
    if review.get("test_only") is not True:
        raise RuntimeError("visual review must be explicitly marked TEST-ONLY")
    if review.get("suite_id") != SUITE_ID:
        raise RuntimeError("visual review suite_id mismatch")
    if review.get("reviewer_type") != "ai":
        raise RuntimeError("this fixture only accepts an explicitly labelled AI review record")
    if review.get("lawyer_approval") is not False or review.get("lawyer_approval_substituted") is not False:
        raise RuntimeError("AI visual review cannot record or substitute lawyer approval")
    expected = {
        (record["scenario_id"], page["page_number"], page["sha256"])
        for record in scenario_records
        for page in record["render"]["pages"]
    }
    actual = {
        (page["scenario_id"], page["page_number"], page["sha256"])
        for page in review.get("pages", [])
    }
    if actual != expected:
        raise RuntimeError("visual review does not cover the exact current page hash set")
    return {
        **review,
        "review_path": relpath(review_path, root),
        "lawyer_approval": False,
        "lawyer_approval_substituted": False,
        "filing_use_allowed": False,
    }


def build(root: Path, *, renderer: Path, skip_render: bool) -> dict[str, Any]:
    root.mkdir(parents=True, exist_ok=True)
    manifest_path = root / "manifest.json"
    previous_manifest = load_json(manifest_path) if manifest_path.is_file() else {}
    previous_by_id = {item["scenario_id"]: item for item in previous_manifest.get("scenarios", [])}

    source, baseline = ensure_source(root)
    source_hash_before = sha256_file(source)
    source_geometry = table_geometry_signature(source)
    scenario_records: list[dict[str, Any]] = []

    for scenario in SCENARIOS:
        scenario_id = scenario["scenario_id"]
        scenario_dir = root / "scenarios" / scenario_id
        output_docx = scenario_dir / "filled.docx"
        derived_bytes = fill_docx_bytes(source, scenario["values"])
        write_bytes_if_changed(output_docx, derived_bytes)
        output_hash = sha256_file(output_docx)
        previous = previous_by_id.get(scenario_id)
        render_dir = scenario_dir / "render"
        reusable = render_is_reusable(previous, output_hash, root)
        if reusable:
            pdf_path = root / previous["render"]["pdf"]["path"]
            pages = [root / item["path"] for item in previous["render"]["pages"]]
            render_action = "reused_hash_verified_render"
        else:
            pdf_path, pages, render_action = render_docx(output_docx, render_dir, renderer, skip_render=skip_render)

        changed_parts_equal, changed_parts = non_document_parts_equal(source, output_docx)
        geometry_equal = table_geometry_signature(output_docx) == source_geometry
        doc_text = document_text(output_docx)
        placeholders_remaining = [marker for marker in PLACEHOLDERS.values() if marker in doc_text]
        actual_fields = field_values(output_docx)
        explicit_blank_ok = all(
            actual_fields.get("可选备注", None) == "" if field == "optional_note" else True
            for field in scenario["explicit_blank_fields"]
        )
        render_inspection = inspect_render(pdf_path, pages, scenario["values"])

        structural_checks = {
            "non_document_ooxml_parts_unchanged": {
                "status": "pass" if changed_parts_equal else "fail",
                "changed_parts": changed_parts,
            },
            "table_geometry_unchanged": {
                "status": "pass" if geometry_equal else "fail",
                "method": "tblW/tblInd/tblLayout/tblGrid/tcW/gridSpan/vMerge and row/cell counts match the immutable source",
            },
            "all_placeholders_resolved": {
                "status": "pass" if not placeholders_remaining else "fail",
                "remaining": placeholders_remaining,
            },
            "explicit_blank_preserved": {
                "status": "pass" if explicit_blank_ok else "fail",
                "fields": scenario["explicit_blank_fields"],
            },
        }
        table_drift_check = {
            "status": "pass" if geometry_equal and changed_parts_equal else "fail",
            "method": "OOXML geometry signature equals the immutable source and every package part except word/document.xml is byte-identical",
        }
        all_machine = (
            all(item["status"] == "pass" for item in structural_checks.values())
            and table_drift_check["status"] == "pass"
            and render_inspection["all_render_checks_passed"]
        )
        scenario_records.append(
            {
                "scenario_id": scenario_id,
                "label": scenario["label"],
                "test_only": True,
                "input_values": scenario["values"],
                "explicit_blank_fields": scenario["explicit_blank_fields"],
                "output_docx": {"path": relpath(output_docx, root), "sha256": output_hash},
                "render": {
                    "renderer": "win-docx-render:LibreOffice+pypdfium2@144dpi",
                    "renderer_script": "win-docx-render:render_docx_win.py",
                    "action": render_action,
                    "pdf": {"path": relpath(pdf_path, root), "sha256": sha256_file(pdf_path)},
                    "page_count": len(pages),
                    "pages": [
                        {
                            "page_number": index,
                            "path": relpath(path, root),
                            "sha256": sha256_file(path),
                        }
                        for index, path in enumerate(pages, start=1)
                    ],
                },
                "structural_checks": structural_checks,
                "machine_checks": {
                    **render_inspection["checks"],
                    "no_table_drift": table_drift_check,
                },
                "page_dimensions": render_inspection["page_dimensions"],
                "all_machine_checks_passed": all_machine,
            }
        )

    source_hash_after = sha256_file(source)
    source_unchanged = source_hash_before == source_hash_after == baseline["sha256"]
    if not source_unchanged:
        raise RuntimeError("immutable source changed during suite generation")
    visual_review = load_visual_review(root, scenario_records)
    manifest = {
        "schema_version": "1.0",
        "suite_id": SUITE_ID,
        "test_only": True,
        "purpose": "Template fill/render boundary regression; never a real-case or filing artifact",
        "generated_by": "tools/build_template_boundary_suite.py",
        "design_contract": {
            "preset": "standard_business_brief",
            "named_overrides": ["Microsoft YaHei East Asian font for Chinese fixture text"],
            "page": "US Letter portrait, one-inch margins",
            "table_width_dxa": 9360,
            "table_indent_dxa": 120,
            "cell_margins_dxa": {"top": 80, "bottom": 80, "start": 120, "end": 120},
        },
        "boundary_contract": {
            "party_name_configured_max_chars": len(MAX_PARTY_NAME),
            "court_name_configured_max_chars": len(MAX_COURT_NAME),
            "large_amount_value": LARGE_AMOUNT,
            "note": "Configured regression boundaries, not statutory or universal business limits.",
        },
        "source": {
            "path": relpath(source, root),
            "sha256": source_hash_after,
            "baseline_path": relpath(root / "source" / "source-baseline.json", root),
            "baseline_sha256": baseline["sha256"],
            "immutable_after_first_build": True,
            "unchanged_during_run": source_unchanged,
            "lawyer_approved_template": False,
        },
        "scenario_count": len(scenario_records),
        "scenarios": scenario_records,
        "all_machine_checks_passed": all(item["all_machine_checks_passed"] for item in scenario_records),
        "visual_review": visual_review,
        "release_gate": {
            "lawyer_visual_approval_required": True,
            "lawyer_visual_approval_present": False,
            "filing_use_allowed": False,
            "template_activation_allowed": False,
            "external_actions_executed": 0,
        },
    }
    write_json_if_changed(manifest_path, manifest)
    return manifest


def main() -> int:
    # Windows consoles may default to a non-UTF-8 codec; CJK output must not crash the builder.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=str(DEFAULT_ROOT), help="Suite output root")
    parser.add_argument("--renderer", help="Explicit win-docx-render helper path")
    parser.add_argument(
        "--skip-render",
        action="store_true",
        help="Do not invoke LibreOffice; require complete existing hash-verified render outputs",
    )
    args = parser.parse_args()
    root = Path(args.root).resolve()
    renderer = find_renderer(args.renderer)
    manifest = build(root, renderer=renderer, skip_render=args.skip_render)
    print(
        json.dumps(
            {
                "ok": manifest["all_machine_checks_passed"],
                "suite_id": manifest["suite_id"],
                "scenario_count": manifest["scenario_count"],
                "source_unchanged": manifest["source"]["unchanged_during_run"],
                "visual_review_status": manifest["visual_review"]["status"],
                "manifest": str(root / "manifest.json"),
            },
            ensure_ascii=False,
        )
    )
    return 0 if manifest["all_machine_checks_passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
